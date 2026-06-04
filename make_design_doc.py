from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document("DESIGN.docx")

# Add a page break then the diagram section at the end
doc.add_page_break()

heading = doc.add_heading("Architecture Diagram", level=1)

caption = doc.add_paragraph("Figure 1 — End-to-end component architecture of the Yann LeCun Digital Twin system.")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.runs[0].italic = True
caption.runs[0].font.size = Pt(10)

image_para = doc.add_paragraph()
image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = image_para.add_run()
run.add_picture(
    r"C:\Users\hp\.gemini\antigravity\brain\46cf450d-75c2-4650-a8a6-54f13c98fb75\architecture_diagram_1780411205995.png",
    width=Inches(6.5)
)

doc.save("DESIGN.docx")
print("Diagram embedded into DESIGN.docx successfully.")
